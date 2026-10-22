
import os
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

from dotenv import load_dotenv
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from supabase import create_client

from telegram.ext import (
    Application,
    CommandHandler,
    CallbackQueryHandler,
    ContextTypes,
)

load_dotenv()

BOT_TOKEN = os.getenv("BOT_TOKEN", "").strip()
ALLOWED_ID = os.getenv("ALLOWED_TELEGRAM_ID", "").strip()
ADMIN_ID = os.getenv("ADMIN_TELEGRAM_ID", "").strip()
EMPLOYEE1_ID = os.getenv("EMPLOYEE1_TELEGRAM_ID", "").strip()
EMPLOYEE2_ID = os.getenv("EMPLOYEE2_TELEGRAM_ID", "").strip()
SUPABASE_URL = os.getenv("SUPABASE_URL", "").strip()
SUPABASE_KEY = os.getenv("SUPABASE_KEY", "").strip()
SUPABASE = create_client(SUPABASE_URL, SUPABASE_KEY) if SUPABASE_URL and SUPABASE_KEY else None
TZ = ZoneInfo("Asia/Tashkent")

MONTHLY_SALARY = 7_000_000  # oylik etalon stavkasi
ADVANCE = 1_950_000
WORK_START_HOUR = 7
LUNCH_SECONDS = 3600
STANDARD_WORK_HOURS_PER_MONTH = 176  # 22 ish kuni × 8 soat
OVERTIME_MULTIPLIER = 1.5

# Foydalanuvchilar keshi; asosiy ma'lumotlar Supabase bazasida saqlanadi.
USERS = {}


def now():
    return datetime.now(TZ)


def fmt(dt):
    return dt.strftime("%d.%m.%Y %H:%M")


def get_user(uid):
    if uid not in USERS:
        USERS[uid] = {"active": None, "history": []}
    return USERS[uid]


def parse_dt(value):
    if not value:
        return None
    dt = datetime.fromisoformat(value.replace("Z", "+00:00"))
    return dt.astimezone(TZ) if dt.tzinfo else dt.replace(tzinfo=TZ)


async def load_user(uid):
    """Xodimning faol seansi va tugallangan ishlarini Supabase'dan yuklaydi."""
    if SUPABASE is None:
        raise RuntimeError(".env faylida SUPABASE_URL va SUPABASE_KEY kerak.")
    response = SUPABASE.table("work_records").select("*").eq("telegram_id", int(uid)).order("start_time").execute()
    user = {"active": None, "history": []}
    for row in (response.data or []):
        start = parse_dt(row.get("start_time"))
        if not start:
            continue
        break_start = parse_dt(row.get("break_start"))
        break_seconds = int(row.get("break_seconds") or (int(row.get("break_minutes") or 0) * 60))
        if row.get("status") == "active":
            user["active"] = {
                "db_id": row["id"], "start": start,
                "break_start": break_start, "break_seconds": break_seconds,
            }
        elif row.get("status") == "completed" and row.get("end_time"):
            end = parse_dt(row["end_time"])
            user["history"].append({
                "db_id": row["id"], "start": start, "end": end,
                "worked_seconds": int(row.get("worked_seconds") or 0),
                "break_seconds": break_seconds,
            })
    USERS[int(uid)] = user
    return user


async def save_user(uid, user):
    """Keshdagi seanslarni ushbu Telegram ID uchun bazaga sinxronlaydi."""
    if SUPABASE is None:
        raise RuntimeError("Supabase ulanishi yo'q. .env qiymatlarini tekshiring.")
    # Har bir yozuvni ID bo'yicha yangilaymiz; eski tarix o'chirilmaydi.
    for rec in user["history"]:
        payload = {
            "telegram_id": int(uid), "employee_name": rec.get("employee_name", ""),
            "work_date": rec["start"].date().isoformat(),
            "start_time": rec["start"].isoformat(), "end_time": rec["end"].isoformat(),
            "break_minutes": int(rec.get("break_seconds", 0) // 60),
            "break_seconds": int(rec.get("break_seconds", 0)),
            "break_start": None, "worked_seconds": int(rec.get("worked_seconds", 0)),
            "status": "completed",
        }
        if rec.get("db_id"):
            SUPABASE.table("work_records").update(payload).eq("id", rec["db_id"]).eq("telegram_id", int(uid)).execute()
        else:
            result = SUPABASE.table("work_records").insert(payload).execute()
            if result.data:
                rec["db_id"] = result.data[0]["id"]
    active = user.get("active")
    if active:
        payload = {
            "telegram_id": int(uid), "employee_name": "",
            "work_date": active["start"].date().isoformat(),
            "start_time": active["start"].isoformat(), "end_time": None,
            "break_minutes": int(active.get("break_seconds", 0) // 60),
            "break_seconds": int(active.get("break_seconds", 0)),
            "break_start": active["break_start"].isoformat() if active.get("break_start") else None,
            "worked_seconds": None, "status": "active",
        }
        if active.get("db_id"):
            SUPABASE.table("work_records").update(payload).eq("id", active["db_id"]).eq("telegram_id", int(uid)).execute()
        else:
            result = SUPABASE.table("work_records").insert(payload).execute()
            if result.data:
                active["db_id"] = result.data[0]["id"]


def duration(seconds):
    seconds = max(0, int(seconds))
    h, rem = divmod(seconds, 3600)
    m = rem // 60
    return f"{h} soat {m} daqiqa"


def keyboard(rows):
    return InlineKeyboardMarkup([
        [InlineKeyboardButton(label, callback_data=data)
         for label, data in row]
        for row in rows
    ])


def main_keyboard(user):
    active = user["active"]

    rows = []

    if not active:
        rows.append([("🟢 Ishni boshlash", "start_work")])
    else:
        rows.append([("🔴 Ishni tugatish", "end_work")])

        if active["break_start"] is None:
            rows.append([("☕ Tanaffusni boshlash", "break_start")])
        else:
            rows.append([("▶️ Tanaffusni tugatish", "break_end")])

    rows.extend([
        [("🕒 Bugun", "today"), ("📅 Haftalik", "week")],
        [("📊 Statistika", "stats"), ("💰 Maoshim", "salary")],
        [("🗂 Ish tarixi", "history"), ("⚙️ Sozlamalar", "settings")],
        [("📥 Excel hisobot", "excel")],
        *([[("👥 Admin paneli", "employees")]] if user.get("is_admin") else []),
        [("ℹ️ Yordam", "help")],
    ])
    return keyboard(rows)


def back_keyboard(back="home"):
    return keyboard([
        [("↩️ Ortga", back), ("🏠 Bosh menyu", "home")]
    ])


def main_text(user, name):
    active = user["active"]
    status = "🟢 Ishda" if active else "⚪ Ish boshlanmagan"

    if active:
        status += f"\n🕒 Boshlangan vaqt: {fmt(active['start'])}"
        if active["break_start"]:
            status += "\n☕ Holat: tanaffusda"
        else:
            status += "\n▶️ Holat: ish davom etmoqda"

    return (
        "━━━━━━━━━━━━━━━━━━\n"
        "       ⏱ WORKTIME\n"
        "   ISH VAQTI NAZORATI\n"
        "━━━━━━━━━━━━━━━━━━\n\n"
        f"👤 Xodim: {name}\n"
        f"📅 Sana: {now().strftime('%d.%m.%Y')}\n"
        f"📍 Holat: {status}\n\n"
        "📋 ASOSIY MENYU\n"
        "Kerakli bo‘limni tanlang 👇"
    )


def completed_seconds(session, end):
    total = int((end - session["start"]).total_seconds())
    breaks = session["break_seconds"]

    if session["break_start"]:
        breaks += int((end - session["break_start"]).total_seconds())

    return max(0, total - breaks)


def user_name(update):
    return update.effective_user.first_name or "foydalanuvchi"


def report_rows(user, since):
    return [
        item for item in user["history"]
        if item["start"] >= since
    ]


def role_for(uid):
    uid = str(uid)
    if ADMIN_ID and uid == ADMIN_ID:
        return "admin"
    if EMPLOYEE1_ID and uid == EMPLOYEE1_ID:
        return "employee1"
    if EMPLOYEE2_ID and uid == EMPLOYEE2_ID:
        return "employee2"
    # Eski bitta-foydalanuvchi sozlamasi bilan vaqtincha moslik.
    if not any((ADMIN_ID, EMPLOYEE1_ID, EMPLOYEE2_ID)) and ALLOWED_ID and uid == ALLOWED_ID:
        return "admin"
    return None


async def allowed(update):
    if role_for(update.effective_user.id):
        return True

    if update.callback_query:
        await update.callback_query.answer(
            "🔒 Sizga ruxsat berilmagan. Admin Telegram ID va xodim IDlarini .env fayliga kiriting.",
            show_alert=True,
        )
    elif update.effective_message:
        await update.effective_message.reply_text(
            "🔒 Bu botdan foydalanishga ruxsatingiz yo‘q."
        )
    return False


def salary_calculation(records):
    """Ishlangan soat bo‘yicha ish haqi va qo‘shimcha haqni hisoblaydi."""
    total_seconds = sum(int(r.get("worked_seconds", 0) or 0) for r in records)
    regular_seconds = sum(min(int(r.get("worked_seconds", 0) or 0), 8 * 3600) for r in records)
    overtime_seconds = sum(max(0, int(r.get("worked_seconds", 0) or 0) - 8 * 3600) for r in records)
    hourly_rate = MONTHLY_SALARY / STANDARD_WORK_HOURS_PER_MONTH
    regular_pay = regular_seconds / 3600 * hourly_rate
    overtime_pay = overtime_seconds / 3600 * hourly_rate * OVERTIME_MULTIPLIER
    gross_pay = regular_pay + overtime_pay
    return {
        "total_seconds": total_seconds,
        "overtime_seconds": overtime_seconds,
        "hourly_rate": hourly_rate,
        "regular_pay": regular_pay,
        "overtime_pay": overtime_pay,
        "gross_pay": gross_pay,
        "advance": ADVANCE,
        "remaining_pay": max(0, gross_pay - ADVANCE),
    }


def money(value):
    return f"{int(round(value)):,}".replace(",", " ") + " so‘m"


def make_excel(records, title):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ish hisoboti"
    ws.append([title])
    ws.append([
        "Sana", "Boshlanish", "Tugash",
        "Tanaffus (daq)", "Ishlangan vaqt (soat)",
        "Kechikish (daq)", "Qo‘shimcha vaqt (soat)", "Qo‘shimcha haq (so‘m)"
    ])

    for rec in records:
        start = rec["start"]
        end = rec["end"]
        late_seconds = max(0, (start.hour * 3600 + start.minute * 60) - WORK_START_HOUR * 3600)
        ws.append([
            start.strftime("%d.%m.%Y"),
            start.strftime("%H:%M"),
            end.strftime("%H:%M"),
            round(rec.get("break_seconds", 0) / 60, 1),
            round(rec.get("worked_seconds", 0) / 3600, 2),
            round(late_seconds / 60, 1),
            round(max(0, rec.get("worked_seconds", 0) - 8 * 3600) / 3600, 2),
            round(max(0, rec.get("worked_seconds", 0) - 8 * 3600) / 3600 * (MONTHLY_SALARY / STANDARD_WORK_HOURS_PER_MONTH) * OVERTIME_MULTIPLIER),
        ])

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=14)
        cell.fill = PatternFill("solid", fgColor="17365D")
    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for col in ws.columns:
        width = max(len(str(cell.value or "")) for cell in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(max(width, 12), 30)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    output.name = "WorkTime_hisobot.xlsx"
    return output


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await allowed(update):
        return

    user = await load_user(update.effective_user.id)
    user["is_admin"] = role_for(update.effective_user.id) == "admin"
    await update.effective_message.reply_text(
        main_text(user, user_name(update)),
        reply_markup=main_keyboard(user),
    )


async def menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()

    if not await allowed(update):
        return

    uid = update.effective_user.id
    user = await load_user(uid)
    user["is_admin"] = role_for(uid) == "admin"
    action = query.data
    current = now()
    message = ""
    back_to = "home"

    # Ishni boshlash
    if action == "start_work":
        if user["active"]:
            message = "⚠️ Ish allaqachon boshlangan."
        else:
            user["active"] = {
                "db_id": None, "start": current,
                "break_start": None, "break_seconds": 0,
            }
            message = (
                "🟢 ISH BOSHLANDI!\n\n"
                f"🕒 Boshlangan vaqt: {fmt(current)}\n"
                "Ish vaqtingiz qayd etildi."
            )

    # Ishni tugatish
    elif action == "end_work":
        active = user["active"]

        if not active:
            message = "⚠️ Avval ishni boshlashingiz kerak."
        elif active["break_start"]:
            message = (
                "☕ Avval tanaffusni tugating, "
                "keyin ishni yakunlang."
            )
        else:
            seconds = completed_seconds(active, current)
            record = {
                "db_id": active.get("db_id"), "start": active["start"],
                "end": current, "worked_seconds": seconds,
                "break_seconds": active["break_seconds"],
            }
            user["history"].append(record)
            user["active"] = None

            message = (
                "🔴 ISH YAKUNLANDI!\n\n"
                f"🟢 Boshlanish: {fmt(record['start'])}\n"
                f"🔴 Tugash: {fmt(current)}\n"
                f"☕ Tanaffus: {duration(record['break_seconds'])}\n"
                f"⏱ Ishlangan vaqt: {duration(seconds)}\n\n"
                "✅ Ish vaqti tarixga saqlandi."
            )

    # Tanaffus boshlash
    elif action == "break_start":
        active = user["active"]

        if not active:
            message = "⚠️ Avval ishni boshlang."
        elif active["break_start"]:
            message = "☕ Tanaffus allaqachon boshlangan."
        else:
            active["break_start"] = current
            message = f"☕ Tanaffus boshlandi: {fmt(current)}"

    # Tanaffus tugatish
    elif action == "break_end":
        active = user["active"]

        if not active:
            message = "⚠️ Faol ish seansi yo‘q."
        elif not active["break_start"]:
            message = "⚠️ Avval tanaffusni boshlang."
        else:
            seconds = int(
                (current - active["break_start"]).total_seconds()
            )
            active["break_seconds"] += max(0, seconds)
            active["break_start"] = None
            message = (
                "▶️ Tanaffus tugadi!\n\n"
                f"☕ Tanaffus davomiyligi: {duration(seconds)}"
            )

    elif action == "home":
        await query.edit_message_text(
            main_text(user, user_name(update)),
            reply_markup=main_keyboard(user),
        )
        return

    elif action == "today":
        back_to = "home"
        today = current.date()
        records = [
            x for x in user["history"]
            if x["start"].date() == today
        ]

        total = sum(x["worked_seconds"] for x in records)
        message = (
            "🕒 BUGUNGI HISOBOT\n"
            "━━━━━━━━━━━━━━━━━━\n\n"
            f"📅 Sana: {today.strftime('%d.%m.%Y')}\n"
            f"📌 Tugallangan ishlar: {len(records)}\n"
            f"⏱ Ishlangan vaqt: {duration(total)}\n"
        )

        active = user["active"]
        if active and active["start"].date() == today:
            if active["break_start"]:
                live_break = int(
                    (current - active["break_start"]).total_seconds()
                )
            else:
                live_break = 0

            live = completed_seconds(active, current)
            message += (
                "\n🟢 Hozir faol ish bor.\n"
                f"Boshlanish: {fmt(active['start'])}\n"
                f"Joriy hisob: {duration(live)}\n"
            )

    elif action == "week":
        back_to = "home"
        since = current - timedelta(days=7)
        records = report_rows(user, since)
        total = sum(x["worked_seconds"] for x in records)

        lines = [
            "📅 HAFTALIK HISOBOT",
            "━━━━━━━━━━━━━━━━━━",
            "",
        ]

        for item in records[-7:]:
            lines.append(
                f"• {item['start'].strftime('%d.%m')} — "
                f"{duration(item['worked_seconds'])}"
            )

        if not records:
            lines.append("Hozircha tugallangan ishlar yo‘q.")

        lines.extend([
            "",
            f"⏱ Jami: {duration(total)}",
        ])
        message = "\n".join(lines)

    elif action == "stats":
        back_to = "home"
        records = user["history"]
        total = sum(x["worked_seconds"] for x in records)
        late = sum(
            1 for x in records
            if x["start"].hour > WORK_START_HOUR
            or (
                x["start"].hour == WORK_START_HOUR
                and x["start"].minute > 0
            )
        )
        extra = sum(
            max(0, x["worked_seconds"] - 8 * 3600)
            for x in records
        )

        message = (
            "📊 UMUMIY STATISTIKA\n"
            "━━━━━━━━━━━━━━━━━━\n\n"
            f"📆 Ish seanslari: {len(records)}\n"
            f"⏱ Jami vaqt: {duration(total)}\n"
            f"⌛ Kech boshlangan seanslar: {late}\n"
            f"➕ 8 soatdan ortiq vaqt: {duration(extra)}"
        )

    elif action == "salary":
        back_to = "home"
        records = user["history"]
        calc = salary_calculation(records)
        days = len({x["start"].date() for x in records})
        regular_pay = calc["gross_pay"] - calc["overtime_pay"]
        message = (
            "💰 MAOSH HISOBOTI\n"
            "━━━━━━━━━━━━━━━━━━\n"
            "📌 1. ISH VAQTI\n"
            f"   📆 Ish kunlari: {days} ta\n"
            f"   ⏱ Jami ishlangan: {duration(calc['total_seconds'])}\n\n"
            "💵 2. HISOB-KITOB\n"
            f"   💳 Soatbay stavka: {money(calc['hourly_rate'])}\n"
            f"   🧮 Oddiy ish haqi: {money(regular_pay)}\n"
            f"   ➕ Qo‘shimcha vaqt: {duration(calc['overtime_seconds'])}\n"
            f"   💸 Qo‘shimcha haq (1.5x): {money(calc['overtime_pay'])}\n"
            f"   ─ Jami hisoblangan: {money(calc['gross_pay'])}\n\n"
            "💳 3. AVANS VA QOLDIQ\n"
            f"   💰 Avans: {money(calc['advance'])}\n"
            f"   ✅ Qoladigan maosh: {money(calc['remaining_pay'])}\n\n"
            "ℹ️ Eslatma\n"
            "• Kechikish faqat hisobotda ko‘rsatiladi.\n"
            "• Kechikish uchun maoshdan pul ushlanmaydi.\n"
            "• Hisob tugallangan ish seanslari asosida.\n"
            "• Soliq va boshqa ushlanmalar hisobga olinmagan."
        )

    elif action == "history":
        back_to = "home"
        records = user["history"][-10:]
        lines = [
            "🗂 OXIRGI ISH TARIXI",
            "━━━━━━━━━━━━━━━━━━",
            "",
        ]

        for item in reversed(records):
            lines.append(
                f"📅 {item['start'].strftime('%d.%m.%Y')}\n"
                f"🕒 {item['start'].strftime('%H:%M')}–"
                f"{item['end'].strftime('%H:%M')}\n"
                f"⏱ {duration(item['worked_seconds'])}\n"
            )

        if not records:
            lines.append("Hozircha tarix bo‘sh.")

        message = "\n".join(lines)

    elif action == "excel":
        back_to = "home"
        records = user["history"]
        if user["active"]:
            message = "⚠️ Avval faol ishni tugating, so‘ng Excel hisobotni oling."
        elif not records:
            message = "📭 Excel yaratish uchun hozircha tugallangan ish yozuvlari yo‘q."
        else:
            file = make_excel(records, f"WorkTime — {user_name(update)}")
            await query.message.reply_document(
                document=file,
                filename="WorkTime_hisobot.xlsx",
                caption="📥 Ish vaqtingiz bo‘yicha Excel hisoboti."
            )
            message = "✅ Excel hisoboti tayyor."

    elif action == "employees":
        back_to = "home"
        if role_for(uid) != "admin":
            message = "🔒 Bu bo‘lim faqat admin uchun."
        else:
            rows = [
                [("📍 Bugungi holat", "admin_today")],
                [("🗂 1-xodim tarixi", "history_employee1"),
                 ("🗂 2-xodim tarixi", "history_employee2")],
                [("📥 1-xodim Excel", "excel_employee1")],
                [("📥 2-xodim Excel", "excel_employee2")],
                [("📥 Adminning Excel hisoboti", "excel")],
            ]
            await query.edit_message_text(
                "👥 ADMIN PANELI\n━━━━━━━━━━━━━━━━━━\n\nKerakli bo‘limni tanlang:",
                reply_markup=keyboard(rows + [[("↩️ Ortga", "home"), ("🏠 Bosh menyu", "home")]])
            )
            return

    elif action == "admin_today":
        back_to = "employees"
        if role_for(uid) != "admin":
            message = "🔒 Bu bo‘lim faqat admin uchun."
        else:
            lines = ["📍 BUGUNGI XODIMLAR HOLATI", "━━━━━━━━━━━━━━━━━━", ""]
            configured = [
                ("1-xodim", EMPLOYEE1_ID),
                ("2-xodim", EMPLOYEE2_ID),
            ]
            for label, target_id in configured:
                if not target_id.isdigit():
                    lines.append(f"👤 {label}: Telegram ID sozlanmagan")
                    continue
                target = await load_user(int(target_id))
                active = target.get("active")
                todays = [r for r in target["history"] if r["start"].date() == current.date()]
                total = sum(r.get("worked_seconds", 0) for r in todays)
                if active:
                    state = f"🟢 Ishda, boshlandi {active['start'].strftime('%H:%M')}"
                    if active.get("break_start"):
                        state += " (tanaffusda)"
                else:
                    state = "⚪ Hozir ish boshlanmagan"
                lines.append(f"👤 {label} — {state}")
                lines.append(f"   Bugun tugallangan: {len(todays)} ta, jami {duration(total)}")
                lines.append("")
            message = "\n".join(lines)

    elif action in ("history_employee1", "history_employee2"):
        back_to = "employees"
        if role_for(uid) != "admin":
            message = "🔒 Bu bo‘lim faqat admin uchun."
        else:
            is_first = action == "history_employee1"
            target_id = EMPLOYEE1_ID if is_first else EMPLOYEE2_ID
            label = "1-xodim" if is_first else "2-xodim"
            target = await load_user(int(target_id)) if target_id.isdigit() else None
            records = target["history"][-10:] if target else []
            lines = [f"🗂 {label.upper()} — OXIRGI ISH TARIXI", "━━━━━━━━━━━━━━━━━━", ""]
            for item in reversed(records):
                lines.append(
                    f"📅 {item['start'].strftime('%d.%m.%Y')}\n"
                    f"🕒 {item['start'].strftime('%H:%M')}–{item['end'].strftime('%H:%M')}\n"
                    f"☕ Tanaffus: {duration(item.get('break_seconds', 0))}\n"
                    f"⏱ Ishlangan: {duration(item.get('worked_seconds', 0))}\n"
                )
            if not records:
                lines.append("Hozircha tugallangan ishlar yo‘q.")
            message = "\n".join(lines)

    elif action in ("excel_employee1", "excel_employee2"):
        back_to = "employees"
        if role_for(uid) != "admin":
            message = "🔒 Bu bo‘lim faqat admin uchun."
        else:
            target_id = EMPLOYEE1_ID if action == "excel_employee1" else EMPLOYEE2_ID
            target_user = await load_user(int(target_id)) if target_id.isdigit() else None
            if not target_user or not target_user["history"]:
                message = "📭 Bu xodim uchun tugallangan ish yozuvlari topilmadi."
            else:
                label = "1-xodim" if action == "excel_employee1" else "2-xodim"
                file = make_excel(target_user["history"], f"WorkTime — {label}")
                await query.message.reply_document(
                    document=file,
                    filename=f"WorkTime_{label}.xlsx",
                    caption=f"📥 {label} ish vaqti hisoboti."
                )
                message = f"✅ {label} Excel hisoboti yuborildi."

    elif action == "settings":
        back_to = "home"
        message = (
            "⚙️ SOZLAMALAR\n"
            "━━━━━━━━━━━━━━━━━━\n"
            "🕒 ISH TARTIBI\n"
            f"   🟢 Ish boshlanishi: {WORK_START_HOUR:02d}:00\n"
            "   ☕ Rejadagi tushlik: 1 soat\n\n"
            "💰 MAOSH SOZLAMALARI\n"
            f"   💵 Oylik maosh: {money(MONTHLY_SALARY)}\n"
            f"   💳 Avans: {money(ADVANCE)}\n"
            f"   🧾 Soatbay stavka: {money(MONTHLY_SALARY / STANDARD_WORK_HOURS_PER_MONTH)}\n"
            "   ➕ Qo‘shimcha soat: 1.5x\n\n"
            "📌 QOIDA\n"
            "   Kechikish maoshdan chegirilmaydi."
        ).replace(",", " ")

    elif action == "help":
        back_to = "home"
        message = (
            "ℹ️ YORDAM\n"
            "━━━━━━━━━━━━━━━━━━\n\n"
            "🟢 Ishni boshlash — vaqtni qayd etadi.\n"
            "🔴 Ishni tugatish — seansni saqlaydi.\n"
            "☕ Tanaffus — tanaffus vaqtini hisoblaydi.\n"
            "📅 Hisobotlar — saqlangan yozuvlarni ko‘rsatadi.\n\n"
            "🏠 Bosh menyu orqali asosiy bo‘limlarga qayting."
        )

    else:
        message = "Bu bo‘lim topilmadi."

    # Har bir amaldan keyin Supabase bilan sinxronlaymiz.
    try:
        await save_user(uid, user)
    except Exception as exc:
        print(f"Supabase saqlash xatosi (Telegram ID {uid}): {exc}")
        message += "\n\n⚠️ Bazaga saqlashda xatolik yuz berdi. Admin terminaldagi xabarni tekshirsin."

    # Ish amallaridan keyin bosh menyudagi holat yangilanadi.
    if action in ("start_work", "end_work", "break_start", "break_end"):
        await query.edit_message_text(
            message,
            reply_markup=keyboard([
                [("🔄 Yangilash", "home")],
                [("↩️ Ortga", "home"), ("🏠 Bosh menyu", "home")],
            ]),
        )
    else:
        await query.edit_message_text(
            message,
            reply_markup=back_keyboard(back_to),
        )


def main():
    if not BOT_TOKEN or "BOTFATHER_TOKENING" in BOT_TOKEN:
        raise RuntimeError(".env faylida haqiqiy BOT_TOKEN qiymatini kiriting.")
    if not SUPABASE_URL or not SUPABASE_KEY:
        raise RuntimeError(".env faylida SUPABASE_URL va SUPABASE_KEY qiymatlarini kiriting.")

    app = Application.builder().token(BOT_TOKEN).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CallbackQueryHandler(menu))

    print("WorkTime bot ishga tushdi!")
    app.run_polling()


if __name__ == "__main__":
    main()
